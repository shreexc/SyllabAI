from hashlib import sha256
from pathlib import Path

from django.conf import settings
from django.core.files.base import ContentFile
from docx import Document
from pypdf import PdfReader

from ..models import LearningChunk, LearningResource

CHUNK_SIZE = 1400
CHUNK_OVERLAP = 180


def extract_pages(resource: LearningResource) -> tuple[list[tuple[str, int | None]], int | None]:
    suffix = Path(resource.file.name).suffix.lower()
    with resource.file.open("rb") as stream:
        if suffix == ".pdf":
            reader = PdfReader(stream, strict=True)
            pages = []
            total_characters = 0
            for number, page in enumerate(reader.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                total_characters += len(page_text)
                if total_characters > settings.LEARNING_MAX_EXTRACTED_CHARACTERS:
                    raise ValueError("Extracted document text exceeds the configured character limit.")
                pages.append((page_text, number))
            return pages, len(reader.pages)
        if suffix == ".txt":
            text = stream.read(settings.LEARNING_MAX_EXTRACTED_CHARACTERS + 1).decode("utf-8-sig").replace("\x00", "").strip()
            if len(text) > settings.LEARNING_MAX_EXTRACTED_CHARACTERS:
                raise ValueError("Extracted document text exceeds the configured character limit.")
            return [(text, None)], None
        if suffix == ".docx":
            document = Document(stream)
            content = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                content.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            text = "\n".join(content)
            if len(text) > settings.LEARNING_MAX_EXTRACTED_CHARACTERS:
                raise ValueError("Extracted document text exceeds the configured character limit.")
            return [(text, None)], None
        return [], None


def extract_text(resource: LearningResource) -> tuple[str, int | None]:
    pages, page_count = extract_pages(resource)
    return "\n\n".join(text for text, _ in pages if text), page_count


def split_text(text: str) -> list[str]:
    normalized = " ".join(text.replace("\x00", "").split())
    chunks: list[str] = []
    cursor = 0
    while cursor < len(normalized):
        end = min(cursor + CHUNK_SIZE, len(normalized))
        if end < len(normalized):
            boundary = normalized.rfind(" ", cursor + CHUNK_SIZE // 2, end)
            if boundary > cursor:
                end = boundary
        chunk = normalized[cursor:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(normalized):
            break
        cursor = max(cursor + 1, end - CHUNK_OVERLAP)
    return chunks


def process_resource(resource_id: str) -> None:
    resource = LearningResource.objects.get(id=resource_id)
    if resource.status == LearningResource.Status.READY:
        return
    resource.status = LearningResource.Status.PROCESSING
    resource.processing_error = ""
    resource.save(update_fields=["status", "processing_error", "updated_at"])
    try:
        pages, page_count = extract_pages(resource)
        text = "\n\n".join(page_text for page_text, _ in pages if page_text)
        if resource.resource_type not in {LearningResource.ResourceType.IMAGE} and len(text) < 30:
            raise ValueError("Not enough readable text was found. Try a text-based PDF, TXT or DOCX document.")
        chunks = [(chunk, page_number) for page_text, page_number in pages for chunk in split_text(page_text)]
        if len(chunks) > 1500:
            raise ValueError("Document is too large after text extraction.")
        LearningChunk.objects.filter(resource=resource, resource_version=resource.resource_version).delete()
        LearningChunk.objects.bulk_create([
            LearningChunk(
                resource=resource,
                resource_version=resource.resource_version,
                index=index,
                text=chunk,
                page_number=page_number,
                content_hash=sha256(chunk.encode("utf-8")).hexdigest(),
            )
            for index, (chunk, page_number) in enumerate(chunks)
        ])
        if resource.processed_text_file:
            resource.processed_text_file.delete(save=False)
        if text:
            resource.processed_text_file.save(
                f"v{resource.resource_version}.txt",
                ContentFile(text.encode("utf-8")),
                save=False,
            )
        resource.metadata = {**resource.metadata, "page_count": page_count, "character_count": len(text), "chunk_count": len(chunks)}
        resource.status = LearningResource.Status.READY
        resource.processing_error = ""
        resource.save(update_fields=["processed_text_file", "metadata", "status", "processing_error", "updated_at"])
    except Exception as error:
        resource.status = LearningResource.Status.FAILED
        resource.processing_error = str(error)[:1000]
        resource.save(update_fields=["status", "processing_error", "updated_at"])
        raise
